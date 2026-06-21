#!/usr/bin/env python3
"""Cisco MyCase PDF Export Assistant — local Playwright utility."""

from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

DEFAULT_LOGIN_URL = "https://mycase.cloudapps.cisco.com/"
DEFAULT_CASE_URL_TEMPLATE = "https://mycase.cloudapps.cisco.com/case/{case_number}"
MANIFEST_NAME = "export_manifest.json"
LOG_NAME = "export.log"
CASE_LOAD_WAIT_MS = 10_000
EXPORT_RETRY_COUNT = 5
EXPORT_RETRY_DELAY_MS = 2_000
VISIBLE_TEXT_PREVIEW_CHARS = 3_000

CASE_COLUMN_NAMES = (
    "case number",
    "case #",
    "case id",
    "sr number",
    "case",
)

MENU_TRIGGER_TEXTS = (
    "Actions",
    "More",
    "Action",
)

EXPORT_TEXT_LABELS = (
    "Save As PDF",
    "Save as PDF",
    "Save PDF",
    "Download PDF",
    "Print",
    "Export",
)

EXPORT_BUTTON_SELECTORS = (
    "button:has-text('Save As PDF')",
    "button:has-text('Save as PDF')",
    "a:has-text('Save As PDF')",
    "a:has-text('Save as PDF')",
    "[role='menuitem']:has-text('Save As PDF')",
    "[role='menuitem']:has-text('Save as PDF')",
    "button:has-text('Save PDF')",
    "button:has-text('Download PDF')",
    "button:has-text('Print')",
    "button:has-text('Export')",
    "a:has-text('Export')",
    "a:has-text('PDF')",
    "[data-testid*='export' i]",
    "[aria-label*='export' i]",
    "[aria-label*='save as pdf' i]",
)

CONTROL_SELECTOR = "button, a, [role='menuitem'], [role='button'], [role='link']"
MENU_TRIGGER_SELECTORS = (
    "button:has-text('Actions')",
    "button:has-text('More')",
    "button:has-text('Action')",
    "[aria-label*='actions' i]",
    "[aria-label*='more' i]",
)

CASE_NUMBER_PATTERN = re.compile(r"^[A-Za-z0-9-]+$")
InputFunc = Callable[[str], str]


@dataclass(frozen=True)
class ExportManifestEntry:
    """One case export result."""

    case_number: str
    status: str
    pdf_path: str | None = None
    html_path: str | None = None
    screenshot_path: str | None = None
    error: str | None = None
    debug_dump: dict[str, Any] | None = None
    updated_at: str = field(default_factory=lambda: _utc_now())

    def to_dict(self) -> dict[str, Any]:
        payload: dict[str, Any] = {
            "case_number": self.case_number,
            "status": self.status,
            "updated_at": self.updated_at,
        }
        if self.pdf_path is not None:
            payload["pdf_path"] = self.pdf_path
        if self.html_path is not None:
            payload["html_path"] = self.html_path
        if self.screenshot_path is not None:
            payload["screenshot_path"] = self.screenshot_path
        if self.error is not None:
            payload["error"] = self.error
        if self.debug_dump is not None:
            payload["debug_dump"] = self.debug_dump
        return payload

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> ExportManifestEntry:
        return cls(
            case_number=str(data["case_number"]),
            status=str(data["status"]),
            pdf_path=data.get("pdf_path"),
            html_path=data.get("html_path"),
            screenshot_path=data.get("screenshot_path"),
            error=data.get("error"),
            debug_dump=data.get("debug_dump"),
            updated_at=str(data.get("updated_at") or _utc_now()),
        )


@dataclass
class ExportManifest:
    """Resume and audit log for case exports."""

    case_list: str
    output_dir: str
    entries: dict[str, ExportManifestEntry] = field(default_factory=dict)
    case_url_template: str | None = None
    created_at: str = field(default_factory=lambda: _utc_now())
    updated_at: str = field(default_factory=lambda: _utc_now())

    def to_dict(self) -> dict[str, Any]:
        payload: dict[str, Any] = {
            "case_list": self.case_list,
            "output_dir": self.output_dir,
            "created_at": self.created_at,
            "updated_at": self.updated_at,
            "entries": {
                case_number: entry.to_dict()
                for case_number, entry in sorted(self.entries.items())
            },
        }
        if self.case_url_template is not None:
            payload["case_url_template"] = self.case_url_template
        return payload

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> ExportManifest:
        entries = {
            case_number: ExportManifestEntry.from_dict(entry)
            for case_number, entry in data.get("entries", {}).items()
        }
        return cls(
            case_list=str(data.get("case_list") or ""),
            output_dir=str(data.get("output_dir") or ""),
            entries=entries,
            case_url_template=data.get("case_url_template"),
            created_at=str(data.get("created_at") or _utc_now()),
            updated_at=str(data.get("updated_at") or _utc_now()),
        )


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize_case_number(value: Any) -> str | None:
    """Normalize a raw case number from spreadsheet cells."""
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in {"none", "nan"}:
        return None
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    text = text.replace(" ", "")
    if not CASE_NUMBER_PATTERN.match(text):
        return None
    return text


def detect_case_column(headers: Sequence[str]) -> str | None:
    """Return the header name that contains Cisco case numbers."""
    lowered = {header: header.strip().lower() for header in headers}
    for candidate in CASE_COLUMN_NAMES:
        for header, header_lower in lowered.items():
            if header_lower == candidate:
                return header
    for header, header_lower in lowered.items():
        if "case" in header_lower:
            return header
    return None


def parse_csv_case_list(path: Path) -> list[str]:
    """Parse case numbers from a CSV export."""
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError(f"No headers found in CSV: {path}")
        column = detect_case_column(reader.fieldnames)
        if column is None:
            raise ValueError(f"Could not detect case number column in CSV: {path}")
        case_numbers: list[str] = []
        seen: set[str] = set()
        for row in reader:
            case_number = normalize_case_number(row.get(column))
            if case_number is None or case_number in seen:
                continue
            seen.add(case_number)
            case_numbers.append(case_number)
        return case_numbers


def parse_xlsx_case_list(path: Path) -> list[str]:
    """Parse case numbers from an Excel export."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required to read Excel case lists. "
            "Install with: pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as exc:
        raise ValueError(f"Excel file is empty: {path}") from exc

    column_index = None
    column_name = detect_case_column(headers)
    if column_name is not None:
        column_index = headers.index(column_name)

    case_numbers: list[str] = []
    seen: set[str] = set()
    for row in rows:
        if column_index is None:
            for value in row:
                case_number = normalize_case_number(value)
                if case_number is not None and case_number not in seen:
                    seen.add(case_number)
                    case_numbers.append(case_number)
                    break
            continue
        if column_index >= len(row):
            continue
        case_number = normalize_case_number(row[column_index])
        if case_number is None or case_number in seen:
            continue
        seen.add(case_number)
        case_numbers.append(case_number)
    return case_numbers


def parse_case_list(path: Path) -> list[str]:
    """Parse case numbers from a Cisco exported case list."""
    if not path.exists():
        raise FileNotFoundError(f"Case list not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return parse_csv_case_list(path)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return parse_xlsx_case_list(path)
    raise ValueError(f"Unsupported case list format: {path.suffix}. Use .csv or .xlsx")


def manifest_path(output_dir: Path) -> Path:
    return output_dir / MANIFEST_NAME


def load_manifest(path: Path) -> ExportManifest | None:
    """Load an existing export manifest if present."""
    if not path.exists():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    return ExportManifest.from_dict(data)


def save_manifest(manifest: ExportManifest, path: Path) -> None:
    """Persist the export manifest."""
    path.parent.mkdir(parents=True, exist_ok=True)
    manifest.updated_at = _utc_now()
    path.write_text(json.dumps(manifest.to_dict(), indent=2) + "\n", encoding="utf-8")


def build_manifest(
    *,
    case_list: Path,
    output_dir: Path,
    existing: ExportManifest | None = None,
) -> ExportManifest:
    """Create or refresh manifest metadata."""
    if existing is None:
        return ExportManifest(
            case_list=str(case_list.resolve()),
            output_dir=str(output_dir.resolve()),
        )
    existing.case_list = str(case_list.resolve())
    existing.output_dir = str(output_dir.resolve())
    return existing


def pdf_path_for_case(output_dir: Path, case_number: str) -> Path:
    return output_dir / f"{case_number}.pdf"


def html_path_for_case(output_dir: Path, case_number: str) -> Path:
    return output_dir / f"{case_number}.html"


def screenshot_path_for_case(output_dir: Path, case_number: str) -> Path:
    return output_dir / f"{case_number}.png"


def should_skip_case(
    case_number: str,
    output_dir: Path,
    manifest: ExportManifest,
    *,
    resume: bool,
) -> bool:
    """Return True when a case should be skipped during resume."""
    if not resume:
        return False
    pdf_path = pdf_path_for_case(output_dir, case_number)
    if pdf_path.exists():
        return True
    entry = manifest.entries.get(case_number)
    return entry is not None and entry.status == "success"


def record_manifest_entry(manifest: ExportManifest, entry: ExportManifestEntry) -> None:
    manifest.entries[entry.case_number] = entry


def configure_logging(output_dir: Path) -> logging.Logger:
    output_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("cisco_mycase_export")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()

    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")

    stream_handler = logging.StreamHandler(sys.stdout)
    stream_handler.setFormatter(formatter)
    logger.addHandler(stream_handler)

    file_handler = logging.FileHandler(output_dir / LOG_NAME, encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    return logger


def build_case_url(case_number: str, template: str) -> str:
    return template.format(case_number=case_number)


def infer_case_url_template(sample_url: str, case_number: str) -> str:
    """Replace a discovered case number in a URL with a template placeholder."""
    if case_number not in sample_url:
        raise ValueError(
            f"Case number {case_number} not found in URL {sample_url!r}. "
            "Open the case detail page before pressing Enter."
        )
    return sample_url.replace(case_number, "{case_number}", 1)


def discover_case_url_template(page: Any, case_number: str) -> str:
    """Discover a reusable case URL template from the current page and frames."""
    candidate_urls = [page.url, *[frame.url for frame in page.frames]]
    for url in candidate_urls:
        if case_number in url:
            return infer_case_url_template(url, case_number)
    raise ValueError(
        f"Could not discover case URL template for {case_number}. "
        "Ensure the open page URL contains the case number."
    )


def normalize_control_text(text: str) -> str:
    """Normalize visible control text for case-insensitive comparison."""
    return " ".join(text.split()).strip().lower()


def matches_export_label(text: str) -> bool:
    """Return True when visible text matches a supported export label."""
    normalized = normalize_control_text(text)
    if not normalized:
        return False
    export_labels = {normalize_control_text(label) for label in EXPORT_TEXT_LABELS}
    return normalized in export_labels


def matches_menu_trigger(text: str) -> bool:
    """Return True when visible text matches an Actions/More menu trigger."""
    normalized = normalize_control_text(text)
    if not normalized:
        return False
    for label in MENU_TRIGGER_TEXTS:
        trigger = normalize_control_text(label)
        if normalized == trigger or normalized.startswith(f"{trigger} "):
            return True
    return False


def _control_text(control: Any) -> str:
    for method_name in ("inner_text", "text_content"):
        method = getattr(control, method_name, None)
        if method is None:
            continue
        try:
            text = method()
        except TypeError:
            text = method(timeout=0)
        if text:
            return str(text).strip()
    return ""


def iter_search_contexts(page: Any) -> list[Any]:
    """Return the page and all child frames to search for controls."""
    contexts: list[Any] = [page]
    for frame in page.frames:
        if frame is page.main_frame:
            continue
        contexts.append(frame)
    return contexts


def iter_visible_controls(context: Any) -> list[tuple[Any, str]]:
    """Return visible controls and text within one page or frame context."""
    controls: list[tuple[Any, str]] = []
    locator = context.locator(CONTROL_SELECTOR)
    try:
        count = locator.count()
    except Exception:
        return controls
    for index in range(count):
        control = locator.nth(index)
        try:
            if not control.is_visible():
                continue
            text = _control_text(control)
            if text:
                controls.append((control, text))
        except Exception:
            continue
    return controls


def collect_visible_control_text(page: Any) -> list[str]:
    """Collect visible button/link/menu text across page and frames."""
    texts: list[str] = []
    seen: set[str] = set()
    for context in iter_search_contexts(page):
        for _, text in iter_visible_controls(context):
            normalized = normalize_control_text(text)
            if normalized and normalized not in seen:
                seen.add(normalized)
                texts.append(text)
    return texts


def find_direct_export_control(context: Any) -> Any | None:
    """Return a visible export control within one page or frame context."""
    for control, text in iter_visible_controls(context):
        if matches_export_label(text):
            return control
    for selector in EXPORT_BUTTON_SELECTORS:
        locator = context.locator(selector)
        try:
            if locator.count() > 0 and locator.first.is_visible():
                return locator.first
        except Exception:
            continue
    return None


def find_menu_trigger_control(context: Any) -> Any | None:
    """Return a visible Actions/More menu trigger within one context."""
    for control, text in iter_visible_controls(context):
        if matches_menu_trigger(text):
            return control
    for selector in MENU_TRIGGER_SELECTORS:
        locator = context.locator(selector)
        try:
            if locator.count() > 0 and locator.first.is_visible():
                return locator.first
        except Exception:
            continue
    return None


def find_export_button(page: Any) -> tuple[Any, Any] | None:
    """Return export context and control, opening Actions/More when required."""
    for context in iter_search_contexts(page):
        direct = find_direct_export_control(context)
        if direct is not None:
            return context, direct

        menu_trigger = find_menu_trigger_control(context)
        if menu_trigger is None:
            continue

        menu_trigger.click()
        page.wait_for_timeout(500)
        direct = find_direct_export_control(context)
        if direct is not None:
            return context, direct
    return None


def find_export_button_with_retries(
    page: Any,
    *,
    retry_count: int = EXPORT_RETRY_COUNT,
    retry_delay_ms: int = EXPORT_RETRY_DELAY_MS,
) -> tuple[Any, Any] | None:
    """Retry export control detection across page and iframe contexts."""
    for attempt in range(retry_count):
        result = find_export_button(page)
        if result is not None:
            return result
        if attempt < retry_count - 1:
            page.wait_for_timeout(retry_delay_ms)
    return None


def build_debug_dump(page: Any) -> dict[str, Any]:
    """Capture page, frame, and control debug details for manifest logging."""
    visible_text_preview = ""
    try:
        visible_text_preview = page.inner_text("body")[:VISIBLE_TEXT_PREVIEW_CHARS]
    except Exception as exc:
        visible_text_preview = f"<unavailable: {exc}>"

    title = ""
    try:
        title = page.title()
    except Exception as exc:
        title = f"<unavailable: {exc}>"

    frame_urls = []
    for frame in page.frames:
        try:
            frame_urls.append(frame.url)
        except Exception:
            frame_urls.append("<unavailable>")

    return {
        "url": page.url,
        "title": title,
        "visible_text_preview": visible_text_preview,
        "frame_urls": frame_urls,
        "control_texts": collect_visible_control_text(page),
    }


def build_export_failure_message(debug_dump: dict[str, Any]) -> str:
    """Build a failure message including visible control text."""
    control_texts = debug_dump.get("control_texts") or []
    if control_texts:
        control_summary = "; ".join(control_texts)
    else:
        control_summary = "(none)"
    return (
        "Export control not found; saved HTML and screenshot. "
        f"URL={debug_dump.get('url')} "
        f"Visible controls: {control_summary}"
    )


def wait_for_case_page(
    page: Any,
    *,
    wait_ms: int = CASE_LOAD_WAIT_MS,
) -> None:
    """Wait for case detail content to finish loading."""
    try:
        page.wait_for_load_state("networkidle", timeout=120_000)
    except Exception:
        pass
    page.wait_for_timeout(wait_ms)


def save_failure_artifacts(page: Any, output_dir: Path, case_number: str) -> tuple[Path, Path]:
    html_path = html_path_for_case(output_dir, case_number)
    screenshot_path = screenshot_path_for_case(output_dir, case_number)
    html_path.write_text(page.content(), encoding="utf-8")
    page.screenshot(path=str(screenshot_path), full_page=True)
    return html_path, screenshot_path


def pause_for_debug(
    case_number: str,
    *,
    input_func: InputFunc = input,
) -> None:
    """Keep the browser open until the user confirms continuation."""
    print(
        f"\nExport control not found for case {case_number}.\n"
        "The browser remains open so you can inspect the page, frames, and menus.\n"
        "Press Enter to continue to the next case.\n",
        flush=True,
    )
    input_func("Press Enter to continue...")


def export_case_with_browser(
    page: Any,
    *,
    case_number: str,
    output_dir: Path,
    case_url_template: str,
    logger: logging.Logger,
    debug_pause_on_failure: bool = False,
    input_func: InputFunc = input,
    case_load_wait_ms: int = CASE_LOAD_WAIT_MS,
    export_retry_count: int = EXPORT_RETRY_COUNT,
) -> ExportManifestEntry:
    """Export one case using an authenticated Playwright page."""
    case_url = build_case_url(case_number, case_url_template)
    pdf_path = pdf_path_for_case(output_dir, case_number)
    logger.info("Opening case %s at %s", case_number, case_url)

    try:
        page.goto(case_url, wait_until="domcontentloaded", timeout=120_000)
        wait_for_case_page(page, wait_ms=case_load_wait_ms)
        export_match = find_export_button_with_retries(
            page,
            retry_count=export_retry_count,
        )

        if export_match is None:
            debug_dump = build_debug_dump(page)
            html_path, screenshot_path = save_failure_artifacts(page, output_dir, case_number)
            message = build_export_failure_message(debug_dump)
            logger.warning("Case %s failed: %s", case_number, message)
            logger.warning("Case %s debug dump: %s", case_number, json.dumps(debug_dump))
            if debug_dump.get("control_texts"):
                logger.warning(
                    "Case %s visible button/menu text: %s",
                    case_number,
                    " | ".join(debug_dump["control_texts"]),
                )
            if debug_pause_on_failure:
                pause_for_debug(case_number, input_func=input_func)
            return ExportManifestEntry(
                case_number=case_number,
                status="failed",
                html_path=str(html_path),
                screenshot_path=str(screenshot_path),
                error=message,
                debug_dump=debug_dump,
            )

        _context, export_button = export_match
        with page.expect_download(timeout=120_000) as download_info:
            export_button.click()
        download = download_info.value
        download.save_as(str(pdf_path))
        logger.info("Case %s exported to %s", case_number, pdf_path)
        return ExportManifestEntry(
            case_number=case_number,
            status="success",
            pdf_path=str(pdf_path),
        )
    except Exception as exc:
        debug_dump = build_debug_dump(page)
        html_path = html_path_for_case(output_dir, case_number)
        screenshot_path = screenshot_path_for_case(output_dir, case_number)
        try:
            html_path.write_text(page.content(), encoding="utf-8")
            page.screenshot(path=str(screenshot_path), full_page=True)
        except Exception:
            html_path = None
            screenshot_path = None
        logger.error("Case %s failed: %s", case_number, exc)
        logger.error("Case %s debug dump: %s", case_number, json.dumps(debug_dump))
        if debug_pause_on_failure:
            pause_for_debug(case_number, input_func=input_func)
        return ExportManifestEntry(
            case_number=case_number,
            status="failed",
            html_path=str(html_path) if html_path else None,
            screenshot_path=str(screenshot_path) if screenshot_path else None,
            error=str(exc),
            debug_dump=debug_dump,
        )


def wait_for_manual_login(
    page: Any,
    login_url: str,
    logger: logging.Logger,
    *,
    input_func: InputFunc = input,
) -> None:
    logger.info("Opening Cisco MyCase login page: %s", login_url)
    page.goto(login_url, wait_until="domcontentloaded", timeout=120_000)
    print(
        "\nManual login required.\n"
        "1. Sign in to Cisco MyCase in the opened browser window.\n"
        "2. Complete MFA/SSO if prompted.\n"
        "3. Confirm you can open a support case page.\n"
        "4. Return here and press Enter to start exports.\n",
        flush=True,
    )
    input_func("Press Enter after login is complete...")
    logger.info("Manual login confirmed by user")


def wait_for_manual_first_case(
    page: Any,
    login_url: str,
    first_case_number: str,
    logger: logging.Logger,
    *,
    input_func: InputFunc = input,
) -> str:
    """Let the user open the first case manually and discover the URL template."""
    logger.info("Opening Cisco MyCase home for manual first-case navigation: %s", login_url)
    page.goto(login_url, wait_until="domcontentloaded", timeout=120_000)
    print(
        "\nManual first-case mode.\n"
        f"1. Sign in to Cisco MyCase if needed.\n"
        f"2. Use MyCase search/list UI to open case {first_case_number}.\n"
        "3. Navigate to the page where Save As PDF is visible.\n"
        "4. Return here and press Enter so the script can record the URL pattern.\n",
        flush=True,
    )
    input_func("Press Enter after the first case detail page is ready...")
    wait_for_case_page(page)
    template = discover_case_url_template(page, first_case_number)
    logger.info("Discovered case URL template: %s", template)
    return template


def run_export(
    *,
    case_numbers: Iterable[str],
    output_dir: Path,
    manifest: ExportManifest,
    delay_seconds: float,
    resume: bool,
    login_url: str,
    case_url_template: str,
    profile_dir: Path,
    headless: bool,
    logger: logging.Logger,
    manual_first_case: bool = False,
    debug_pause_on_failure: bool = False,
    input_func: InputFunc = input,
) -> ExportManifest:
    from playwright.sync_api import sync_playwright

    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir.mkdir(parents=True, exist_ok=True)
    case_list = list(case_numbers)
    if not case_list:
        return manifest

    active_template = manifest.case_url_template or case_url_template

    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            user_data_dir=str(profile_dir),
            headless=headless,
            accept_downloads=True,
        )
        page = context.pages[0] if context.pages else context.new_page()

        if manual_first_case:
            active_template = wait_for_manual_first_case(
                page,
                login_url,
                case_list[0],
                logger,
                input_func=input_func,
            )
            manifest.case_url_template = active_template
            save_manifest(manifest, manifest_path(output_dir))
            logger.info("Using discovered URL template for remaining cases")
        else:
            wait_for_manual_login(page, login_url, logger, input_func=input_func)

        for index, case_number in enumerate(case_list, start=1):
            if should_skip_case(case_number, output_dir, manifest, resume=resume):
                logger.info("Skipping case %s (already exported)", case_number)
                continue

            entry = export_case_with_browser(
                page,
                case_number=case_number,
                output_dir=output_dir,
                case_url_template=active_template,
                logger=logger,
                debug_pause_on_failure=debug_pause_on_failure,
                input_func=input_func,
            )
            record_manifest_entry(manifest, entry)
            save_manifest(manifest, manifest_path(output_dir))

            if delay_seconds > 0 and index < len(case_list):
                logger.info("Waiting %.1f seconds before next case", delay_seconds)
                time.sleep(delay_seconds)

        context.close()
    return manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Export Cisco MyCase support cases to PDF using a manual browser login.",
    )
    parser.add_argument("--case-list", required=True, help="Path to Cisco case list Excel/CSV")
    parser.add_argument(
        "--output",
        default="exports/cisco_mycase",
        help="Directory for exported PDFs and manifest",
    )
    parser.add_argument("--delay", type=float, default=5.0, help="Delay in seconds between cases")
    parser.add_argument("--resume", action="store_true", help="Skip already exported cases")
    parser.add_argument(
        "--profile-dir",
        default=".browser_profile",
        help="Persistent Playwright browser profile directory",
    )
    parser.add_argument("--login-url", default=DEFAULT_LOGIN_URL, help="Cisco MyCase login URL")
    parser.add_argument(
        "--case-url-template",
        default=DEFAULT_CASE_URL_TEMPLATE,
        help="Case page URL template containing {case_number}",
    )
    parser.add_argument(
        "--manual-first-case",
        action="store_true",
        help="Open first case manually and discover the working case URL template",
    )
    parser.add_argument(
        "--debug-pause-on-failure",
        action="store_true",
        help="Pause with browser open when export control is not found",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Run browser headless (manual login not supported)",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    case_list = Path(args.case_list)
    output_dir = Path(args.output)
    profile_dir = Path(args.profile_dir)
    logger = configure_logging(output_dir)

    case_numbers = parse_case_list(case_list)
    if not case_numbers:
        logger.error("No case numbers found in %s", case_list)
        return 1

    existing = load_manifest(manifest_path(output_dir))
    manifest = build_manifest(case_list=case_list, output_dir=output_dir, existing=existing)
    if existing and existing.case_url_template:
        manifest.case_url_template = existing.case_url_template
    save_manifest(manifest, manifest_path(output_dir))

    logger.info("Loaded %d case numbers from %s", len(case_numbers), case_list)
    logger.info("Output directory: %s", output_dir.resolve())
    logger.info("Case URL template: %s", args.case_url_template)
    if args.resume:
        logger.info("Resume enabled")
    if args.manual_first_case:
        logger.info("Manual first-case mode enabled")
    if args.debug_pause_on_failure:
        logger.info("Debug pause on failure enabled")

    run_export(
        case_numbers=case_numbers,
        output_dir=output_dir,
        manifest=manifest,
        delay_seconds=args.delay,
        resume=args.resume,
        login_url=args.login_url,
        case_url_template=args.case_url_template,
        profile_dir=profile_dir,
        headless=args.headless,
        logger=logger,
        manual_first_case=args.manual_first_case,
        debug_pause_on_failure=args.debug_pause_on_failure,
    )

    successes = sum(1 for entry in manifest.entries.values() if entry.status == "success")
    failures = sum(1 for entry in manifest.entries.values() if entry.status == "failed")
    logger.info("Export complete. success=%d failed=%d", successes, failures)
    return 0 if failures == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
